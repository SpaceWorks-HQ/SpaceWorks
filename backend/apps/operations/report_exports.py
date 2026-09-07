import csv
from datetime import datetime
from io import BytesIO, StringIO

from django.http import HttpResponse, StreamingHttpResponse
from openpyxl import Workbook

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class StreamingCsvResponse(StreamingHttpResponse):
    """A streamed CSV that still answers `.content` for callers that read the whole body.

    Django's streaming response deliberately has no `content`; the test client and any
    caller that wants the bytes get them here by materialising the stream once.
    """

    @property
    def content(self):
        # The generator is one-shot; cache so a second read sees the same bytes.
        if not hasattr(self, "_materialized_content"):
            self._materialized_content = b"".join(self.streaming_content)
        return self._materialized_content


def _csv_lines(rows, provenance):
    buffer = StringIO()
    writer = csv.writer(buffer)
    # Written raw, not through the writer: the JSON filters contain commas and the whole
    # point of the line is to stay readable as one `# key=value ...` comment.
    yield provenance.header_line() + "\r\n"
    for row in rows:
        writer.writerow([_export_cell(value) for value in row])
        yield buffer.getvalue()
        buffer.seek(0)
        buffer.truncate(0)


def _csv_response(rows, filename, *, provenance):
    response = StreamingCsvResponse(_csv_lines(rows, provenance), content_type="text/csv")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def csv_bytes(rows, *, provenance):
    return "".join(_csv_lines(rows, provenance)).encode("utf-8")


def xlsx_bytes(rows, *, provenance):
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Report")
    for row in rows:
        sheet.append([_xlsx_cell(value) for value in row])
    provenance_sheet = workbook.create_sheet("Provenance")
    for key, value in provenance.items():
        provenance_sheet.append([key, _export_cell(value)])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _xlsx_response(rows, filename, *, provenance):
    response = HttpResponse(xlsx_bytes(rows, provenance=provenance), content_type=XLSX_CONTENT_TYPE)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def _xlsx_cell(value):
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.replace(tzinfo=None)
    return _export_cell(value)


def _export_cell(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return f"'{value}"
    return value
