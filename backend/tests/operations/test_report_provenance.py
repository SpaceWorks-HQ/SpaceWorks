"""Every CSV/XLSX export carries a provenance row/sheet and CSV streams."""

from io import BytesIO

import pytest
from django.http import StreamingHttpResponse
from openpyxl import load_workbook

from apps.accounts.models import User
from apps.operations.report_exports import csv_bytes, xlsx_bytes
from apps.operations.report_exports_provenance import build_provenance
from tests.return_helpers import (
    authenticated_client,
    make_issued_request,
    make_member,
    make_product,
    make_space,
    make_user,
)

pytestmark = pytest.mark.django_db


def _provenance_fields(line):
    assert line.startswith("# generated_at="), line
    pairs = line[2:].split(" filters=", 1)
    fields = dict(item.split("=", 1) for item in pairs[0].split(" "))
    fields["filters"] = pairs[1]
    return fields


def test_csv_export_streams_and_starts_with_the_provenance_row():
    space = make_space("prov-csv")
    manager = make_member("prov-csv-manager", space)
    product = make_product(space, name="Prov Scope")
    make_issued_request(space, manager, [(product, 1)])

    response = authenticated_client(manager).get(
        f"/api/v1/admin/makerspace/{space.id}/reports/most-lent/export"
    )

    assert response.status_code == 200
    assert isinstance(response, StreamingHttpResponse)
    lines = b"".join(response.streaming_content).decode().splitlines()
    fields = _provenance_fields(lines[0])
    assert fields["generated_by"] == manager.username
    assert fields["makerspace_id"] == str(space.id)
    assert fields["report_key"] == "most-lent"
    assert fields["report_version"] == "1"
    assert fields["filters"] == '{"grain":"day"}'
    assert lines[1] == "product_name,times_lent,total_quantity_lent"
    assert lines[2] == "Prov Scope,1,1"


def test_aggregate_and_ledger_exports_carry_provenance_too():
    space = make_space("prov-aggregate")
    manager = make_member("prov-aggregate-manager", space)
    superadmin = make_user(
        "prov-super", role=User.Role.SUPERADMIN, access_status=User.AccessStatus.ACTIVE,
    )

    aggregate = authenticated_client(superadmin).get(
        "/api/v1/admin/reports/most-lent/export?start=2026-01-01&end=2026-12-31"
    )
    ledger = authenticated_client(manager).get(
        f"/api/v1/admin/makerspace/{space.id}/ledger/export?format=csv&overdue=true"
    )

    aggregate_fields = _provenance_fields(aggregate.content.decode().splitlines()[0])
    assert aggregate_fields["makerspace_id"] == "all"
    assert aggregate_fields["generated_by"] == superadmin.username
    assert '"start":"2026-01-01T00:00:00' in aggregate_fields["filters"]
    assert '"end":"2027-01-01T00:00:00' in aggregate_fields["filters"]
    ledger_fields = _provenance_fields(ledger.content.decode().splitlines()[0])
    assert ledger_fields["report_key"] == "ledger"
    assert ledger_fields["makerspace_id"] == str(space.id)
    assert '"overdue":true' in ledger_fields["filters"]


def test_xlsx_export_has_a_provenance_sheet_after_the_report_sheet():
    space = make_space("prov-xlsx")
    manager = make_member("prov-xlsx-manager", space)
    make_product(space, name="Prov Meter", available_quantity=9, damaged_quantity=1, lost_quantity=0)

    response = authenticated_client(manager).get(
        f"/api/v1/admin/makerspace/{space.id}/reports/damaged-lost/export?format=xlsx"
    )

    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == ["Report", "Provenance"]
    assert [cell.value for cell in workbook["Report"][1]] == ["product_name", "damaged_quantity", "lost_quantity"]
    provenance = {row[0]: row[1] for row in workbook["Provenance"].iter_rows(values_only=True)}
    assert provenance["report_key"] == "damaged-lost"
    assert provenance["generated_by"] == manager.username
    assert provenance["makerspace_id"] == str(space.id)
    assert provenance["report_version"] == "1"


def test_byte_renderers_share_the_provenance_with_the_http_helpers():
    provenance = build_provenance(
        "most-lent", version=3, makerspace_id=None, generated_by="schedule:7", filters={"window_days": 7},
    )
    rows = [["product_name", "times_lent"], ["=SUM(A1)", 2]]

    csv_lines = csv_bytes(rows, provenance=provenance).decode().splitlines()
    assert csv_lines[0] == provenance.header_line()
    assert "makerspace_id=all report_key=most-lent report_version=3" in csv_lines[0]
    assert csv_lines[2] == "'=SUM(A1),2"
    workbook = load_workbook(BytesIO(xlsx_bytes(rows, provenance=provenance)))
    assert workbook["Report"]["A2"].value == "'=SUM(A1)"
    assert dict(workbook["Provenance"].iter_rows(values_only=True))["generated_by"] == "schedule:7"
